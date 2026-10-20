"""
clean_contacts.py — Nettoie un fichier de contacts en bazar et sort un Excel propre.

Exemple concret de ce que j'automatise pour mes clients :
un CSV exporté à l'arrache (doublons, majuscules incohérentes, téléphones
dans 5 formats différents, lignes vides) devient un fichier Excel trié,
dédoublonné et formaté, prêt à l'emploi.

Usage :
    python clean_contacts.py                      # génère un exemple + le nettoie (démo)
    python clean_contacts.py contacts.csv         # nettoie votre fichier -> contacts_propre.xlsx
    python clean_contacts.py contacts.csv out.xlsx
    python clean_contacts.py --selfcheck          # lance les tests internes
"""

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLONNES = ["Nom", "Email", "Telephone"]


def clean_text(s: str) -> str:
    """Enlève les espaces superflus et met un nom proprement en Majuscule Initiale."""
    s = " ".join((s or "").split())  # espaces multiples/bords -> un seul
    return s.title()


def clean_email(s: str) -> str:
    return " ".join((s or "").split()).lower()


def clean_phone(s: str) -> str:
    """Normalise un numéro français en '0X XX XX XX XX'. Garde la valeur brute si invalide."""
    digits = re.sub(r"\D", "", s or "")
    if digits.startswith("0033"):                       # 0033 6... -> 06...
        digits = "0" + digits[4:]
    elif digits.startswith("33") and len(digits) == 11:  # +33 6... -> 06...
        digits = "0" + digits[2:]
    if len(digits) != 10 or not digits.startswith("0"):
        return (s or "").strip()  # on ne jette rien : on garde tel quel pour relecture
    return " ".join(digits[i:i + 2] for i in range(0, 10, 2))


def clean_rows(rows: list[dict]) -> list[dict]:
    """Nettoie chaque champ, supprime les lignes vides, dédoublonne par email, trie par nom."""
    vus = set()
    propres = []
    for r in rows:
        nom = clean_text(r.get("Nom", ""))
        email = clean_email(r.get("Email", ""))
        tel = clean_phone(r.get("Telephone", ""))
        if not nom and not email and not tel:
            continue  # ligne totalement vide
        cle = email or nom  # doublon = même email (sinon même nom)
        if cle in vus:
            continue
        vus.add(cle)
        propres.append({"Nom": nom, "Email": email, "Telephone": tel})
    propres.sort(key=lambda r: r["Nom"].lower())
    return propres


def load_rows(csv_path: Path) -> list[dict]:
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_xlsx(rows: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"

    entete_fill = PatternFill("solid", fgColor="2F5496")
    entete_font = Font(color="FFFFFF", bold=True)
    ws.append(COLONNES)
    for col, _ in enumerate(COLONNES, start=1):
        c = ws.cell(row=1, column=col)
        c.fill = entete_fill
        c.font = entete_font
        c.alignment = Alignment(horizontal="left")

    for r in rows:
        ws.append([r[c] for c in COLONNES])

    ws.freeze_panes = "A2"  # entête figée au scroll
    for col, nom in enumerate(COLONNES, start=1):
        largeur = max([len(nom)] + [len(str(r[nom])) for r in rows] + [10])
        ws.column_dimensions[get_column_letter(col)].width = largeur + 4

    wb.save(out_path)


def make_sample(path: Path) -> None:
    """Crée un CSV volontairement en bazar pour la démo."""
    brut = [
        ["Nom", "Email", "Telephone"],
        ["  jean   DUPONT ", "Jean.Dupont@MAIL.com ", "06 12 34 56 78"],
        ["marie martin", "marie.martin@mail.com", "+33 7 11 22 33 44"],
        ["JEAN DUPONT", "jean.dupont@mail.com", "06.12.34.56.78"],  # doublon
        ["", "", ""],                                               # ligne vide
        ["sophie   bernard", "SOPHIE.bernard@mail.com", "0033688776655"],
        ["luc petit ", " luc.petit@mail.com", "0611223344"],
        ["  ", "  ", "  "],                                         # ligne vide
    ]
    with open(path, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(brut)


def nettoyer_fichier(src: Path, out: Path) -> int:
    rows = clean_rows(load_rows(src))
    write_xlsx(rows, out)
    return len(rows)


def _selfcheck() -> None:
    assert clean_text("  jean   DUPONT ") == "Jean Dupont"
    assert clean_email(" A@B.COM ") == "a@b.com"
    assert clean_phone("+33 6 12 34 56 78") == "06 12 34 56 78"
    assert clean_phone("06.12.34.56.78") == "06 12 34 56 78"
    assert clean_phone("0033688776655") == "06 88 77 66 55"
    assert clean_phone("pas un numero") == "pas un numero"  # on garde le brut
    sample = [
        {"Nom": "JEAN DUPONT", "Email": "jean.dupont@mail.com", "Telephone": "06.12.34.56.78"},
        {"Nom": " jean dupont ", "Email": "Jean.Dupont@mail.com", "Telephone": "0612345678"},
        {"Nom": "", "Email": "", "Telephone": ""},
        {"Nom": "anne", "Email": "anne@mail.com", "Telephone": "0711223344"},
    ]
    propres = clean_rows(sample)
    assert len(propres) == 2, propres          # 1 doublon + 1 vide retirés
    assert propres[0]["Nom"] == "Anne"         # trié par nom
    print("OK — tous les tests passent.")


def main(argv: list[str]) -> None:
    if argv and argv[0] == "--selfcheck":
        _selfcheck()
        return

    if not argv:  # mode démo : on fabrique l'exemple puis on le nettoie
        src = Path("exemple_sale.csv")
        make_sample(src)
        out = Path("exemple_propre.xlsx")
        n = nettoyer_fichier(src, out)
        print(f"Démo : {src} -> {out} ({n} contacts propres)")
        return

    src = Path(argv[0])
    out = Path(argv[1]) if len(argv) > 1 else src.with_name(src.stem + "_propre.xlsx")
    n = nettoyer_fichier(src, out)
    print(f"{src} -> {out} ({n} contacts propres)")


if __name__ == "__main__":
    main(sys.argv[1:])
